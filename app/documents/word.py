"""DOCX exchange by stable bookmark fields and project/batch metadata."""
from __future__ import annotations

import json
import os
import sqlite3
import uuid
from datetime import datetime, timezone
from pathlib import Path

from docx import Document
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches

from app.ocr.review import TextBlockService
from app.word_protocol import DISPLAY_LABELS, NOTES, PROTOCOL_VERSION, SOURCE, TEXT_ID, TRANSLATION


FIELDS = (TEXT_ID, SOURCE, TRANSLATION, NOTES)


def timestamp() -> str:
    return datetime.now(timezone.utc).isoformat()


def bookmark_name(field: str, uid: str) -> str:
    return f"{field}__{uid.replace('-', '_')}"


def bookmark(paragraph, name: str, number: int) -> None:
    start = OxmlElement("w:bookmarkStart")
    start.set(qn("w:id"), str(number))
    start.set(qn("w:name"), name)
    paragraph._p.insert(0, start)
    end = OxmlElement("w:bookmarkEnd")
    end.set(qn("w:id"), str(number))
    paragraph._p.append(end)


class WordExchange:
    """Own DOCX exchange records; page order and translated display headings are not IDs."""
    def __init__(self, connection: sqlite3.Connection, project: Path) -> None:
        self.db = connection
        self.project = project

    def _batch(self, batch_id: str):
        batch = self.db.execute("SELECT * FROM batches WHERE id=?", (batch_id,)).fetchone()
        if batch is None:
            raise ValueError("Unknown batch")
        return batch

    def export(self, batch_id: str, mode: str = "light") -> Path:
        if mode not in ("light", "standard", "full"):
            raise ValueError("Unknown Word export mode")
        batch = self._batch(batch_id)
        project = self.db.execute("SELECT id,name FROM projects").fetchone()
        revision = self.db.execute("SELECT COALESCE(MAX(export_revision),0)+1 FROM document_exports WHERE batch_id=?",
                                   (batch_id,)).fetchone()[0]
        pages = self.db.execute("""SELECT p.* FROM batch_pages bp JOIN pages p ON p.id=bp.page_id
            WHERE bp.batch_id=? ORDER BY bp.position""", (batch_id,)).fetchall()
        document = Document()
        document.core_properties.title = f"{project['name']} — Batch {batch['batch_number']}"
        document.core_properties.keywords = json.dumps({"protocol": PROTOCOL_VERSION,
            "project": project["id"], "batch": batch_id, "revision": revision}, separators=(",", ":"))
        document.add_heading(document.core_properties.title, 0)
        serial = 1
        review = TextBlockService(self.db)
        for page in pages:
            blocks = review.for_page(page["id"])
            document.add_heading(f"{page['page_uid']} — {page['label']}", level=1)
            if mode in ("standard", "full"):
                self._add_page_preview(document, dict(page))
            table = document.add_table(rows=1, cols=4)
            table.style = "Table Grid"
            for cell, field in zip(table.rows[0].cells, FIELDS):
                cell.text = DISPLAY_LABELS[field]
            for block in blocks:
                translation = review.translation(block["id"])
                values = (block["text_uid"], block["source_text"],
                          translation["text"] if translation else "", translation["notes"] if translation else "")
                cells = table.add_row().cells
                for cell, field, value in zip(cells, FIELDS, values):
                    cell.text = value
                    bookmark(cell.paragraphs[0], bookmark_name(field, block["text_uid"]), serial)
                    serial += 1
                if mode == "full":
                    self._add_block_crop(cells[1], dict(page), dict(block))
        output = self.project / "documents" / f"Batch_{batch['batch_number']:03d}_r{revision}_{mode}.docx"
        output.parent.mkdir(parents=True, exist_ok=True)
        temporary = output.with_name(f"{uuid.uuid4().hex[:12]}.docx")
        try:
            document.save(temporary)
            os.replace(temporary, output)
            with self.db:
                self.db.execute("""INSERT INTO document_exports
                    (id,project_id,batch_id,protocol_version,export_revision,path,created_at)
                    VALUES(?,?,?,?,?,?,?)""", (str(uuid.uuid4()), project["id"], batch_id,
                    PROTOCOL_VERSION, revision, str(output), timestamp()))
            return output
        finally:
            temporary.unlink(missing_ok=True)

    def _page_image(self, page: dict):
        from app.pages.page_service import PageService
        from app.images.image_loader import ImageLoader
        from app.pdf.render import PdfRenderService
        current = next(p for p in PageService(self.db, self.project).list_pages() if p["id"] == page["id"])
        if current["source_kind"] == "pdf":
            return PdfRenderService(self.project).render(current)
        return ImageLoader().load(Path(current["path"]), 900)

    @staticmethod
    def _image_stream(image, max_width: int = 700):
        import io
        from PySide6.QtCore import QBuffer, QIODevice, Qt
        if image.width() > max_width:
            image = image.scaledToWidth(max_width, Qt.TransformationMode.SmoothTransformation)
        buffer = QBuffer()
        buffer.open(QIODevice.OpenModeFlag.WriteOnly)
        if not image.save(buffer, "JPEG", 72):
            raise OSError("Cannot compress Word preview")
        return io.BytesIO(bytes(buffer.data()))

    def _add_page_preview(self, document, page: dict) -> None:
        image = self._page_image(page)
        document.add_picture(self._image_stream(image), width=Inches(4.5))

    def _add_block_crop(self, cell, page: dict, block: dict) -> None:
        import json as _json
        from PySide6.QtCore import QRect
        image = self._page_image(page)
        x0, y0, x1, y1 = _json.loads(block["bbox_json"])
        sx, sy = image.width() / page["width"], image.height() / page["height"]
        pad = 15
        box = QRect(max(0, round(x0*sx)-pad), max(0, round(y0*sy)-pad),
                    round((x1-x0)*sx)+2*pad, round((y1-y0)*sy)+2*pad)
        cropped = image.copy(box.intersected(image.rect()))
        if not cropped.isNull():
            cell.add_paragraph().add_run().add_picture(self._image_stream(cropped, 320), width=Inches(1.5))

    def import_document(self, path: Path) -> dict:
        document = Document(path)
        try:
            metadata = json.loads(document.core_properties.keywords)
        except (ValueError, TypeError) as exc:
            raise ValueError("Missing Links Manga Workspace document protocol") from exc
        project_id = self.db.execute("SELECT id FROM projects").fetchone()[0]
        if metadata.get("protocol") != PROTOCOL_VERSION or metadata.get("project") != project_id:
            raise ValueError("Document protocol or project UUID does not match")
        batch_id = metadata.get("batch")
        self._batch(batch_id)
        export = self.db.execute("""SELECT id FROM document_exports WHERE batch_id=? AND export_revision=?
            AND protocol_version=?""", (batch_id, metadata.get("revision"), PROTOCOL_VERSION)).fetchone()
        if export is None:
            raise ValueError("Unknown export revision")
        known = {row["text_uid"]: row for row in self.db.execute("""SELECT t.* FROM text_blocks t
            JOIN batch_pages bp ON bp.page_id=t.page_id WHERE bp.batch_id=?""", (batch_id,))}
        found: dict[str, list[str]] = {}
        for table in document.tables:
            for row in table.rows:
                for cell in row.cells:
                    for tag in cell._tc.xpath(".//w:bookmarkStart"):
                        name = tag.get(qn("w:name"), "")
                        prefix = TRANSLATION + "__"
                        if name.startswith(prefix):
                            uid = name[len(prefix):].replace("_", "-")
                            found.setdefault(uid, []).append("\n".join(p.text for p in cell.paragraphs).strip())
        report = {key: 0 for key in ("matched", "imported", "empty", "unknown", "duplicate", "conflict", "missing")}
        report["missing"] = len(set(known) - set(found))
        review = TextBlockService(self.db)
        for uid, values in found.items():
            if uid not in known:
                report["unknown"] += 1
                continue
            if len(values) > 1:
                report["duplicate"] += 1
                continue
            report["matched"] += 1
            value = values[0]
            if not value:
                report["empty"] += 1
                continue
            current = review.translation(known[uid]["id"])
            if current and current["text"].strip() and current["text"] != value:
                report["conflict"] += 1
                continue
            review.save_translation(known[uid]["id"], value)
            report["imported"] += 1
        with self.db:
            self.db.execute("INSERT INTO document_imports(id,document_export_id,path,report_json,created_at) VALUES(?,?,?,?,?)",
                            (str(uuid.uuid4()), export["id"], str(path), json.dumps(report), timestamp()))
        return report
