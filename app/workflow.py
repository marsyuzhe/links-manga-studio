"""Small transactional workflow operations over stable page and text identifiers."""
from __future__ import annotations

import json
import uuid
from datetime import datetime, timezone
from pathlib import Path

from app.history import HistoryService
from app.ocr.review import TextBlockService


class WorkflowService:
    def __init__(self, connection, project: Path) -> None:
        self.db = connection
        self.project = project

    def erase_blocks(self, block_ids: list[str], mode: str = "fill") -> int:
        if mode not in ("fill", "telea", "ns"):
            raise ValueError("Invalid erase mode")
        ids = list(dict.fromkeys(block_ids))
        if not ids:
            return 0
        history = HistoryService(self.db)
        group_id = str(uuid.uuid4())
        with self.db:
            for block_id in ids:
                before = history.snapshot("text_blocks", block_id)
                if before is None or not before["active"]:
                    raise ValueError(f"Unknown active text block: {block_id}")
                self.db.execute("UPDATE text_blocks SET erase_data_json=?,erase_status='erased',revision=revision+1 WHERE id=?",
                                (json.dumps({"mode": mode, "color": "#ffffff", "radius": 3}), block_id))
                history.record("text_blocks", block_id, before, "erase", group_id)
        return len(ids)

    def fill_translation(self, block_id: str, text: str) -> None:
        row = self.db.execute("SELECT active FROM text_blocks WHERE id=?", (block_id,)).fetchone()
        if row is None or not row["active"]:
            raise ValueError("Text block no longer exists")
        if not text.strip():
            raise ValueError("Please enter a translation first")
        review = TextBlockService(self.db)
        previous = review.translation(block_id)
        if not previous or previous["text"] != text:
            review.save_translation(block_id, text, previous["notes"] if previous else "")

    def restore_erase(self, block_id: str) -> None:
        history = HistoryService(self.db)
        before = history.snapshot("text_blocks", block_id)
        if before is None or not before["active"]:
            raise ValueError("Text block no longer exists")
        with self.db:
            self.db.execute("UPDATE text_blocks SET erase_data_json='{}',erase_status='none',revision=revision+1 WHERE id=?",
                            (block_id,))
            history.record("text_blocks", block_id, before, "restore_erase")

    def erase_and_refill(self, block_id: str, text: str, mode: str = "fill") -> None:
        if not text.strip():
            raise ValueError("Please enter a translation first")
        if mode not in ("fill", "telea", "ns"):
            raise ValueError("Invalid erase mode")
        history = HistoryService(self.db)
        before_block = history.snapshot("text_blocks", block_id)
        if before_block is None or not before_block["active"]:
            raise ValueError("Text block no longer exists")
        previous = TextBlockService(self.db).translation(block_id)
        translation_id = previous["id"] if previous else str(uuid.uuid4())
        before_translation = history.snapshot("translations", translation_id)
        group_id = str(uuid.uuid4())
        with self.db:
            self.db.execute("""UPDATE text_blocks SET erase_data_json=?,erase_status='erased',
                typeset_status='ready',revision=revision+1 WHERE id=?""",
                (json.dumps({"mode": mode, "color": "#ffffff", "radius": 3}), block_id))
            history.record("text_blocks", block_id, before_block, "erase_and_refill", group_id)
            timestamp = datetime.now(timezone.utc).isoformat()
            if previous:
                status = previous["status"] if previous["text"] == text else "human_edited"
                self.db.execute("""UPDATE translations SET text=?,status=?,revision=revision+1,
                    updated_at=? WHERE id=?""", (text, status, timestamp, translation_id))
            else:
                self.db.execute("""INSERT INTO translations
                    (id,text_block_id,language,text,notes,status,updated_at,created_at) VALUES(?,?,'zh_CN',?,'','human_edited',?,?)""",
                    (translation_id, block_id, text, timestamp, timestamp))
            history.record("translations", translation_id, before_translation, "erase_and_refill", group_id)

    def export_source(self, page_ids: list[str], destination: Path) -> int:
        ids = list(dict.fromkeys(page_ids))
        if not ids:
            raise ValueError("No pages selected")
        placeholders = ",".join("?" for _ in ids)
        rows = self.db.execute(f"""SELECT p.id AS page_id,p.page_uid,p.display_order,b.text_uid,b.reading_order,b.source_text
            FROM pages p LEFT JOIN text_blocks b ON b.page_id=p.id AND b.active=1
            WHERE p.id IN ({placeholders}) AND p.deleted_at IS NULL
            ORDER BY p.display_order,b.reading_order,b.sequence""", ids).fetchall()
        if len({row["page_id"] for row in rows}) != len(ids):
            raise ValueError("Unknown page in source export")
        lines = ["Page UID\tDisplay Order\tText UID\tReading Order\tOriginal Text"]
        for row in rows:
            if row["text_uid"]:
                source = (row["source_text"] or "").replace("\t", " ").replace("\r", " ").replace("\n", " ")
                lines.append(f"{row['page_uid']}\t{row['display_order']}\t{row['text_uid']}\t{row['reading_order']}\t{source}")
        destination.parent.mkdir(parents=True, exist_ok=True)
        temporary = destination.with_suffix(destination.suffix + ".tmp")
        try:
            temporary.write_text("\n".join(lines) + "\n", encoding="utf-8-sig")
            temporary.replace(destination)
        finally:
            temporary.unlink(missing_ok=True)
        return len(lines) - 1

    def export_source_docx(self, page_ids: list[str], destination: Path) -> int:
        """Export original text with stable IDs for human translation handoff."""
        from docx import Document
        ids = list(dict.fromkeys(page_ids))
        if not ids:
            raise ValueError("No pages selected")
        placeholders = ",".join("?" for _ in ids)
        rows = self.db.execute(f"""SELECT p.id AS page_id,p.page_uid,p.display_order,
            b.text_uid,b.reading_order,b.source_text FROM pages p
            LEFT JOIN text_blocks b ON b.page_id=p.id AND b.active=1
            WHERE p.id IN ({placeholders}) AND p.deleted_at IS NULL
            ORDER BY p.display_order,b.reading_order,b.sequence""", ids).fetchall()
        if len({row["page_id"] for row in rows}) != len(ids):
            raise ValueError("Unknown page in source export")
        document = Document()
        document.add_heading("Original Text", 0)
        last_page = None
        count = 0
        for row in rows:
            if row["page_id"] != last_page:
                document.add_heading(f"Page {row['display_order']} · {row['page_uid']}", level=1)
                last_page = row["page_id"]
            if row["text_uid"]:
                document.add_heading(row["text_uid"], level=2)
                document.add_paragraph(row["source_text"] or "")
                count += 1
        destination.parent.mkdir(parents=True, exist_ok=True)
        temporary = destination.with_suffix(".tmp.docx")
        try:
            document.save(temporary)
            temporary.replace(destination)
        finally:
            temporary.unlink(missing_ok=True)
        return count
