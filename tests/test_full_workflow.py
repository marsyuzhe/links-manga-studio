import json
import uuid

from docx import Document

from app.batches.service import BatchService
from app.documents.word import WordExchange
from app.pages.page_service import PageService
from app.pdf.importer import PdfImporter
from app.rendering.export import export_pages
from test_pdf import make_pdf


def test_thousand_page_end_to_end(qapp, project_factory, tmp_path):
    project = project_factory(name="千页全流程")
    pdf = make_pdf(tmp_path / "漫画.pdf")
    assert PdfImporter(project.path).import_file(pdf) == 1000
    pages = PageService(project.connection, project.path).list_pages()
    batches = BatchService(project.connection)
    batch_ids = batches.create_for_unassigned(100)
    assert len(batch_ids) == 10
    # Lightweight OCR task simulation. A separate test runs the real ONNX engine.
    with project.connection:
        project.connection.executemany("""INSERT INTO text_blocks
            (id,page_id,text_uid,sequence,reading_order,bbox_json,source_text,ocr_confidence)
            VALUES(?,?,?,?,?,?,?,?)""", [(str(uuid.uuid4()), page["id"],
            page["page_uid"] + "-T001", 1, 1, json.dumps([20, 20, 180, 100]),
            f"SOURCE {i}", .9) for i, page in enumerate(pages, 1)])
    word = WordExchange(project.connection, project.path)
    exported = word.export(batch_ids[0], "light")
    document = Document(exported)
    assert len(document.tables[0].rows) == 2
    document.tables[0].rows[1].cells[2].paragraphs[0].add_run("译文一")
    document.tables[1].rows[1].cells[2].paragraphs[0].add_run("译文二")
    translated = tmp_path / "returned.docx"
    document.save(translated)
    report = word.import_document(translated)
    assert report["matched"] == 100 and report["imported"] == 2 and report["empty"] == 98
    assert project.connection.execute("SELECT COUNT(*) FROM translations WHERE text!=''").fetchone()[0] == 2
    result = export_pages(project.path, [page["id"] for page in pages])
    assert result["status"] == "completed" and result["completed"] == 1000
    output = project.path / "exports" / f"render_{result['task_id'][:8]}"
    assert len(list(output.glob("*.png"))) == 1000
    assert (output / "P0001.png").is_file() and (output / "P1000.png").is_file()
