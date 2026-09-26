from docx import Document

from app.batches.service import BatchService
from app.documents.word import WordExchange
from app.importers.folder_importer import scan_files
from app.importers.image_importer import ImageImporter
from app.ocr.pipeline import OCRPipeline
from app.ocr.review import TextBlockService
from app.pages.page_service import PageService
from conftest import make_image
from test_ocr import FakeBackend


def test_word_stable_bookmark_import_after_row_move(project_factory, tmp_path):
    project = project_factory()
    paths = [make_image(tmp_path / f"{i}.png", 300, 200, i) for i in range(2)]
    importer = ImageImporter(project.path)
    importer.commit(importer.validate(scan_files(paths)))
    pages = PageService(project.connection, project.path).list_pages()
    for page in pages:
        OCRPipeline(project.connection, project.path, FakeBackend()).process_page(page)
    batch_id = BatchService(project.connection).create_for_unassigned()[0]
    word = WordExchange(project.connection, project.path)
    exported = word.export(batch_id, "light")
    document = Document(exported)
    first, second = document.tables
    first.rows[1].cells[2].text = "谢谢"
    second.rows[1].cells[2].text = "你好"
    # Setting text removes bookmarks: restore by editing paragraph run only.
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn
    from app.documents.word import bookmark, bookmark_name
    for table, uid, text in ((first, "P0001-T001", "谢谢"), (second, "P0002-T001", "你好")):
        cell = table.rows[1].cells[2]
        bookmark(cell.paragraphs[0], bookmark_name("LMW_TRANSLATION", uid), 100 + int(uid[1:5]))
    document.element.body.remove(second._element)
    document.element.body.insert(0, second._element)
    edited = tmp_path / "translated.docx"
    document.save(edited)
    report = word.import_document(edited)
    assert report == {"matched": 2, "imported": 2, "empty": 0, "unknown": 0,
                      "duplicate": 0, "conflict": 0, "missing": 0}
    review = TextBlockService(project.connection)
    assert review.translation(review.for_page(pages[0]["id"])[0]["id"])["text"] == "谢谢"
    assert review.translation(review.for_page(pages[1]["id"])[0]["id"])["text"] == "你好"
    assert word.import_document(edited)["imported"] == 2


def test_word_standard_full_compression(project_factory, tmp_path):
    project = project_factory()
    path = make_image(tmp_path / "large.png", 900, 1350)
    importer = ImageImporter(project.path)
    importer.commit(importer.validate(scan_files([path])))
    page = PageService(project.connection, project.path).list_pages()[0]
    OCRPipeline(project.connection, project.path, FakeBackend()).process_page(page)
    batch = BatchService(project.connection).create_for_unassigned()[0]
    word = WordExchange(project.connection, project.path)
    standard = word.export(batch, "standard")
    full = word.export(batch, "full")
    assert standard.stat().st_size < 200_000
    assert full.stat().st_size < 300_000
    assert len(Document(standard).inline_shapes) == 1
    assert len(Document(full).inline_shapes) == 2
